#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
[AgentDesk] 生成「官方 10 大类 → 41 子类」权威映射（Go 源码）

为什么要用脚本生成而不是手写：
    初版按「业务直觉」手工分组，结果 41 条里 33 条大类归错 ——
    大类准确率只有 16.9%，而子类准确率是 67.7%。落差全部来自这个错位。
    生成可以保证「与官方数据逐字一致」，杜绝人为归类偏差。

用法:
    python scripts/gen_scene_major_map.py
"""

from __future__ import annotations

import io
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

XLSX = Path("D:/浏览器下载/赛题 1：数据共情者-业务数据.xlsx")
OUT = Path("agent-desk/internal/services/cs_scene_major_map.go")

COL_SCENE_MAJOR = 8
COL_SCENE_MINOR = 9


def build_mapping() -> dict[str, set[str]]:
    wb = load_workbook(XLSX, data_only=True, read_only=True)
    ws = wb["聊天记录"]
    rows = ws.iter_rows(values_only=True)
    next(rows)

    m: dict[str, set[str]] = defaultdict(set)
    for r in rows:
        mj, mi = r[COL_SCENE_MAJOR], r[COL_SCENE_MINOR]
        if mj and mi:
            m[str(mj).strip()].add(str(mi).strip())
    return m


def render(m: dict[str, set[str]]) -> str:
    total = sum(len(v) for v in m.values())
    assert len(m) == 10, f"官方应为 10 大类，实际 {len(m)}"
    assert total == 41, f"官方应为 41 子类，实际 {total}"

    o = io.StringIO()
    o.write("package services\n\n")
    o.write("// ============================================================================\n")
    o.write("// [AgentDesk-BEGIN] 官方 10 大类 → 41 子类 权威映射\n")
    o.write("//\n")
    o.write("// ★ 本表由 scripts/gen_scene_major_map.py 从官方 xlsx 直接生成，请勿手改。\n")
    o.write("//   重新生成：python scripts/gen_scene_major_map.py\n")
    o.write("//\n")
    o.write("// 为什么必须生成而非手写（血泪教训）：\n")
    o.write("//   初版按「业务直觉」自行分组，41 条里 33 条大类归错 ——\n")
    o.write("//     物流停滞疑似丢件 → 误归「物流」，官方是「物流异常」\n")
    o.write("//     色号选择        → 误归「咨询」，官方是「产品咨询」\n")
    o.write("//     发票申请        → 误归「咨询」，官方是「订单服务」\n")
    o.write("//   后果：子类判对、大类却报出官方不存在的名字，\n")
    o.write("//         大类准确率 16.9% vs 子类准确率 67.7%（落差全在此）。\n")
    o.write("//\n")
    o.write("// 用法：所有**对外输出**的大类必须经 sceneMajorOf() 查这张表，\n")
    o.write("//       不要直接读 SceneCategoryRule.SceneMajor（那只是调试字段）。\n")
    o.write("// ============================================================================\n\n")

    o.write("// sceneMinorToMajor 41 个子类到官方大类的唯一真相源。\n")
    o.write("var sceneMinorToMajor = map[string]string{\n")
    for mj in sorted(m, key=lambda k: (-len(m[k]), k)):
        o.write(f"\t// {mj}（{len(m[mj])} 类）\n")
        for s in sorted(m[mj]):
            o.write(f'\t"{s}": "{mj}",\n')
        o.write("\n")
    o.write("}\n\n")

    o.write("// officialSceneMajors 官方 10 大类，按该大类下的子类数降序。\n")
    o.write("var officialSceneMajors = []string{\n")
    for mj in sorted(m, key=lambda k: (-len(m[k]), k)):
        o.write(f'\t"{mj}",\n')
    o.write("}\n\n")
    o.write("// [AgentDesk-END]\n")
    return o.getvalue()


def main() -> int:
    if not XLSX.exists():
        print(f"[FAIL] 找不到官方数据文件: {XLSX}")
        return 1

    m = build_mapping()
    src = render(m)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    with io.open(OUT, "w", encoding="utf-8", newline="\n") as f:
        f.write(src)

    print(f"[OK] 已生成 {OUT}")
    print(f"     10 大类 / 41 子类")
    for mj in sorted(m, key=lambda k: (-len(m[k]), k)):
        print(f"     {mj:<8} {len(m[mj])} 类")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
